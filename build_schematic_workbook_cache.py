from __future__ import annotations

import hashlib
import json
import re
import warnings
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


APP_DIR = Path(__file__).resolve().parent
WORKBOOK = APP_DIR.parents[1] / "Data Nations" / "well schematic Rokan.xlsx"
CACHE_DIR = APP_DIR / "data" / "schematic_workbook_cache"


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_text(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize(value: object) -> str:
    return re.sub(r"[^A-Z0-9]+", "", clean_text(value).upper())


def orientation(value: object) -> str:
    match = re.search(r"([VDH])(?:1)?$", normalize(value))
    return match.group(1) if match else ""


def depth_interval(value: object) -> tuple[float, float] | None:
    match = re.search(
        r"(-?\d+(?:\.\d+)?)\s*(?:-|\u2013|\u2014|TO)\s*(-?\d+(?:\.\d+)?)",
        clean_text(value).replace(",", ""),
        flags=re.IGNORECASE,
    )
    if match is None:
        return None
    first, second = float(match.group(1)), float(match.group(2))
    return min(first, second), max(first, second)


def json_value(value: object) -> object:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def parse_sheet(worksheet: Any) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    rows = list(worksheet.iter_rows(values_only=True))
    header_index = -1
    columns: dict[str, int] = {}
    for row_index, row in enumerate(rows[:15]):
        labels = {clean_text(value).upper(): index for index, value in enumerate(row)}
        if "WELL DIAGRAM" in labels and "CASING" in labels:
            header_index = row_index
            columns = labels
            break
    if header_index < 0:
        return [], []

    formation_column = columns.get("CASING SHOE")
    interval_column = columns.get("WELL DIAGRAM")
    hole_column = columns.get("HOLE DETAILS")
    casing_column = columns.get("CASING")
    md_column = columns.get("CASING DEPTH (MD)")
    tvd_column = columns.get("CASING DEPTH (TVD)")
    # Completion status (Open/Closed/NZBP) is stored in the workbook's
    # ``Hole Details`` column on the same row as the interval.  The previous
    # parser looked one column to the left and therefore lost the status.
    status_column = hole_column

    def value_at(row: tuple[Any, ...], column: int | None) -> Any:
        return row[column] if column is not None and column < len(row) else None

    casing_rows: list[dict[str, object]] = []
    interval_rows: list[dict[str, object]] = []
    for source_row, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        casing_text = clean_text(value_at(row, casing_column))
        hole_text = clean_text(value_at(row, hole_column))
        md_value = pd.to_numeric(value_at(row, md_column), errors="coerce")
        tvd_value = pd.to_numeric(value_at(row, tvd_column), errors="coerce")
        # Hole details also contains completion-state labels.  Do not count
        # those rows as casing strings unless casing or depth evidence exists.
        if casing_text or pd.notna(md_value) or pd.notna(tvd_value):
            casing_rows.append(
                {
                    "Hole Details": hole_text,
                    "Casing": casing_text,
                    "Top (ft)": 0.0,
                    "Base (ft)": json_value(md_value),
                    "Casing Depth (MD)": json_value(md_value),
                    "Casing Depth (TVD)": json_value(tvd_value),
                    "Source Row": source_row,
                }
            )
        interval_text = clean_text(value_at(row, interval_column))
        parsed = depth_interval(interval_text)
        if parsed is not None:
            formation = clean_text(value_at(row, formation_column))
            status = clean_text(value_at(row, status_column))
            interval_rows.append(
                {
                    "Formation": formation,
                    "Sand": formation,
                    "Interval": interval_text,
                    "Top (ft)": parsed[0],
                    "Base (ft)": parsed[1],
                    "Status": status,
                    "Kegiatan": "Perforation interval",
                    "Temuan": f"{formation} {status}".strip(),
                    "Depth (feet)": interval_text,
                    "Source Row": source_row,
                }
            )
    return casing_rows, interval_rows


def main() -> None:
    if not WORKBOOK.exists():
        raise FileNotFoundError(WORKBOOK)
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    stat = WORKBOOK.stat()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        workbook = load_workbook(
            WORKBOOK,
            read_only=False,
            data_only=True,
            keep_links=False,
        )
    entries: list[dict[str, object]] = []
    templates = {"VERTICAL", "DIRECTIONAL", "HORIZONTAL"}
    try:
        for position, worksheet in enumerate(workbook.worksheets):
            if clean_text(worksheet.title).upper() in templates:
                continue
            casing, intervals = parse_sheet(worksheet)
            slug = re.sub(r"[^a-z0-9]+", "_", worksheet.title.lower()).strip("_")
            stem = f"{position:02d}_{slug}"
            data_name = f"{stem}.json"
            image_name = ""
            image_format = ""
            images = list(getattr(worksheet, "_images", []))
            if images:
                largest = max(
                    images,
                    key=lambda image: float(getattr(image, "width", 0))
                    * float(getattr(image, "height", 0)),
                )
                try:
                    image_format = str(getattr(largest, "format", "png") or "png").lower()
                    image_name = f"{stem}.{image_format}"
                    (CACHE_DIR / image_name).write_bytes(largest._data())
                except Exception:
                    image_name = ""
                    image_format = ""
            payload = {
                "sheet_name": worksheet.title,
                "orientation": orientation(worksheet.title),
                "image_format": image_format,
                "sheet_render_file": f"{stem}_sheet.png",
                "casing": casing,
                "intervals": intervals,
            }
            (CACHE_DIR / data_name).write_text(
                json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            entries.append(
                {
                    "sheet_name": worksheet.title,
                    "orientation": orientation(worksheet.title),
                    "casing_count": len(casing),
                    "interval_count": len(intervals),
                    "image_available": bool(image_name),
                    "data_file": data_name,
                    "image_file": image_name,
                    "sheet_render_file": f"{stem}_sheet.png",
                }
            )
    finally:
        workbook.close()
    index = {
        "source_path": str(WORKBOOK),
        "source_size": int(stat.st_size),
        "source_mtime_ns": int(stat.st_mtime_ns),
        "source_sha256": file_sha256(WORKBOOK),
        "entries": entries,
    }
    (CACHE_DIR / "index.json").write_text(
        json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"SCHEMATIC_CACHE_OK sheets={len(entries)} path={CACHE_DIR}")


if __name__ == "__main__":
    main()
